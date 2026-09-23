from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from datetime import UTC, date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import Engine, text

from ust_pipeline.repository import Repository

from .common import (
    DemoSafetyError,
    ValidationFailure,
    assert_safe_database_state,
    catalog_objects,
    create_demo_engine,
    demo_database_url,
    migration_version,
    repository_root,
    server_info,
)
from .loader import MANIFEST_PATH, demo_validation, object_counts, verify_manifest


SPECIAL_SHEETS = [
    "00_README",
    "01_TABLE_INDEX",
    "02_VIEW_INDEX",
    "03_COLUMN_COVERAGE",
    "04_LINEAGE_TRACE",
    "05_VALIDATION",
]
FORBIDDEN_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
FORMULA_PREFIXES = ("=", "+", "-", "@")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="000000")
THIN_BLUE = Side(style="thin", color="9EADBD")


TABLE_PURPOSES = {
    "ingestion_run": "수집 실행 상태·건수·체크포인트 감사",
    "source_request": "HTTP 요청·응답 상태 감사",
    "source_snapshot": "원문 URL·해시·raw 저장 경로",
    "data_quality_issue": "검증·review·NULL 사유 기록",
    "security_master": "CUSIP 기준 증권 마스터",
    "auction_event": "입찰 공고·일정의 현재 상태",
    "auction_result": "입찰 결과와 상품별 Stop",
    "auction_bidder_allocation": "참여자별 응찰·낙찰 금액",
    "auction_revision": "입찰 A→B→A 정정 이력",
    "normalized_record_lineage": "정규화 레코드에서 raw snapshot까지 lineage",
    "qra_refunding": "분기 Refunding 발표 단위",
    "qra_document": "QRA 문서의 안정 식별자",
    "qra_document_version": "QRA 내용·parser·격리 버전 이력",
    "qra_borrowing_estimate": "Treasury 차입 실적·전망",
    "qra_auction_size": "Treasury 입찰 규모와 TBAC 권고",
    "qra_supply": "Gross·Maturing·Net 명목 쿠폰 공급",
    "qra_financing_mix": "분기 조달 mix와 bills 잔여",
    "qra_tga_path": "TGA 실제·가정 앵커",
    "qra_guidance": "상품군별 공식 issuance guidance",
    "qra_tentative_auction": "잠정 입찰 일정",
    "qra_buyback_operation": "잠정 buyback 일정",
    "qra_buyback_policy": "buyback 용도별 상한",
    "qra_dealer_survey": "Primary Dealer survey 문서·대상기간",
    "qra_dealer_survey_value": "Dealer expected 값과 통계량",
    "derived_metric_method": "비공식 UI proxy 계산 방법 버전",
    "derived_metric_method_parameter": "ui_proxy_v1 tenor weight",
    "derived_metric_result": "방법·입력·산식이 고정된 계산 결과",
}

VIEW_METADATA = {
    "v_auction_prior_comparable": ("입찰 KPI/직전 비교", "직전 유효 결과와 최근 6개 유효 Stop 평균", "유효 Stop이 없으면 NULL, n 별도"),
    "v_auction_dashboard": ("입찰 대시보드", "Stop·응찰배수×100·배정비중·Other 잔여", "0 분모·결과 미공고는 NULL"),
    "v_qra_supply_monitor": ("QRA 공급", "Gross−Maturing=Net, Net×tenor weight", "공식 공급과 DERIVED_UI_PROXY 분리"),
    "v_qra_comparison": ("QRA 최신/직전", "현재·직전 전망과 명시 delta", "직전 문서 미수집 시 cross-document 값 NULL"),
    "v_qra_tga_path": ("QRA TGA", "실제 시작·분기말·중간 peak 앵커", "일일 구간을 보간하지 않음"),
    "v_qra_dealer_outlook": ("Dealer outlook", "expected/current−1", "current 근거가 없으면 current와 변화율 NULL"),
    "v_qra_auction_size_comparison": ("QRA 입찰 규모", "같은 월·상품·만기의 직전 Refunding 비교", "비교행이 없으면 prior/delta NULL"),
    "v_ingestion_status": ("수집 상태", "source별 최근 시도·성공·수신·거절", "성공 이력이 없으면 last_success NULL"),
}


def safe_sheet_name(prefix: str, object_name: str, used: set[str]) -> str:
    base = FORBIDDEN_SHEET_CHARS.sub("_", f"{prefix}_{object_name}").strip("'") or "Sheet"
    if len(base) > 31:
        digest = hashlib.sha1(base.encode("utf-8")).hexdigest()[:6]
        base = f"{base[:24]}_{digest}"
    candidate = base
    counter = 2
    while candidate.casefold() in {item.casefold() for item in used}:
        suffix = f"_{counter}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate)
    return candidate


def safe_source_text(value: str) -> str:
    return f"'{value}" if value.startswith(FORMULA_PREFIXES) else value


def stable_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value
    if isinstance(value, time):
        return value
    if isinstance(value, (dict, list, tuple, set)):
        return safe_source_text(stable_json(value))
    if isinstance(value, memoryview):
        return value.tobytes().hex()
    if isinstance(value, bytes):
        return value.hex()
    if isinstance(value, str):
        return safe_source_text(value)
    return value


def parse_db_spec(root: Path) -> tuple[dict[str, str], dict[tuple[str, str], dict[str, str]]]:
    text_value = (root / "docs" / "DB_SPEC.md").read_text(encoding="utf-8")
    purposes: dict[str, str] = {}
    columns: dict[tuple[str, str], dict[str, str]] = {}
    current: str | None = None
    for line in text_value.splitlines():
        section = re.fullmatch(r"## ([a-z][a-z0-9_]*)", line.strip())
        if section:
            current = section.group(1)
            continue
        if current and current not in purposes and line and not line.startswith(("|", "`", "#")):
            purposes[current] = line.split("。", 1)[0].strip()
        if current and line.startswith("| `"):
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if len(cells) >= 7:
                column = cells[0].strip("`")
                columns[(current, column)] = {
                    "description": cells[1], "type": cells[2].strip("`"), "unit": cells[3],
                    "nullable": cells[4], "source": cells[6],
                }
    return purposes, columns


def column_catalog(engine: Engine, object_name: str) -> list[dict[str, Any]]:
    with engine.connect() as connection:
        rows = connection.execute(text("""
            SELECT a.attname, pg_catalog.format_type(a.atttypid,a.atttypmod),
                   NOT a.attnotnull AS nullable, a.attnum
            FROM pg_attribute a
            JOIN pg_class c ON c.oid=a.attrelid
            JOIN pg_namespace n ON n.oid=c.relnamespace
            WHERE n.nspname='ust' AND c.relname=:name AND a.attnum>0 AND NOT a.attisdropped
            ORDER BY a.attnum
        """), {"name": object_name}).all()
    return [{"name": row[0], "type": row[1], "nullable": bool(row[2]), "ordinal": row[3]} for row in rows]


def constraint_metadata(engine: Engine, table: str) -> tuple[list[str], list[str]]:
    with engine.connect() as connection:
        rows = connection.execute(text("""
            SELECT contype, pg_get_constraintdef(oid)
            FROM pg_constraint
            WHERE conrelid=to_regclass(:qualified) AND contype IN ('p','u')
            ORDER BY contype, conname
        """), {"qualified": f"ust.{table}"}).all()
        pk_rows = connection.execute(text("""
            SELECT a.attname
            FROM pg_index i
            JOIN unnest(i.indkey) WITH ORDINALITY k(attnum,ord) ON true
            JOIN pg_attribute a ON a.attrelid=i.indrelid AND a.attnum=k.attnum
            WHERE i.indrelid=to_regclass(:qualified) AND i.indisprimary
            ORDER BY k.ord
        """), {"qualified": f"ust.{table}"}).all()
    uniques = [definition for kind, definition in rows if kind == "u"]
    return [row[0] for row in pk_rows], uniques


def update_policy(table: str) -> tuple[str, str]:
    if table in {"ingestion_run", "source_request", "source_snapshot", "data_quality_issue", "auction_revision", "normalized_record_lineage", "qra_document_version", "derived_metric_result"}:
        return "APPEND/AUDIT", "raw·실행·정정·버전·계산 이력 보존"
    if table in {"derived_metric_method", "derived_metric_method_parameter"}:
        return "METHOD VERSION", "방법 버전을 추가하고 당시 입력 고정"
    return "BUSINESS-KEY UPSERT", "현재행 upsert, 변경 이력은 revision/version·raw에 보존"


def _query_rows(engine: Engine, object_name: str, pk_columns: list[str] | None = None) -> tuple[list[str], list[tuple[Any, ...]]]:
    columns = [item["name"] for item in column_catalog(engine, object_name)]
    order = ""
    if pk_columns:
        order = " ORDER BY " + ",".join(f'"{item}"' for item in pk_columns)
    with engine.connect() as connection:
        rows = connection.execute(text(f'SELECT * FROM ust."{object_name}"{order}')).all()
    if not pk_columns:
        rows = sorted(rows, key=lambda row: stable_json([excel_value(item) for item in row]))
    return columns, [tuple(row) for row in rows]


def _apply_header(ws, columns: list[str], metadata: list[dict[str, Any]] | None = None) -> None:
    for index, name in enumerate(columns, start=1):
        cell = ws.cell(1, index, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_BLUE)
        if metadata:
            item = metadata[index - 1]
            cell.comment = Comment(
                "\n".join([
                    item.get("description", name),
                    f"타입: {item.get('type', '')}",
                    f"단위: {item.get('unit', '—')}",
                    f"NULL: {'허용' if item.get('nullable') else '불가'}",
                    f"원천: {item.get('source', '')}",
                ]),
                "User",
            )
    ws.row_dimensions[1].height = 34


def _finish_sheet(ws, *, freeze: str = "A2", filter_rows: bool = True) -> None:
    ws.freeze_panes = freeze
    if filter_rows and ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if isinstance(cell.value, date) and not isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(cell.value, time):
                cell.number_format = "hh:mm:ss"
    for index in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, index).value or "") for row in range(1, min(ws.max_row, 200) + 1)]
        width = min(45, max(10, max((len(item) for item in values), default=0) + 2))
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_matrix(ws, headers: list[str], rows: list[list[Any]], metadata: list[dict[str, Any]] | None = None) -> None:
    _apply_header(ws, headers, metadata)
    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            ws.cell(row_index, column_index, excel_value(value))
    _finish_sheet(ws)


def _column_metadata_for_object(
    engine: Engine,
    object_name: str,
    spec_columns: dict[tuple[str, str], dict[str, str]],
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for column in column_catalog(engine, object_name):
        documented = spec_columns.get((object_name, column["name"]), {})
        result.append({
            **column,
            "description": documented.get("description", column["name"].replace("_", " ")),
            "unit": documented.get("unit", "—"),
            "source": documented.get("source", "DB view 계산" if object_name.startswith("v_") else "DB catalog/정규화"),
        })
    return result


def _all_null_reason(table: str, column: str) -> str:
    if table == "qra_dealer_survey_value" and column in {"current_auction_size_usd", "calculated_pct_change"}:
        return "공식 survey가 expected 값만 제공. current와 delta는 근거가 없어 NULL"
    if table == "qra_dealer_survey" and column in {"survey_as_of_date", "expected_increase_timing"}:
        return "공식 문서에 일/증액 시점이 없어 MONTH precision 또는 NULL 유지"
    if column in {"etag", "last_modified"}:
        return "offline fixture 응답에 조건부 요청 헤더가 없음"
    if "error" in column:
        return "성공 실행/요청에는 오류가 없음"
    if table.startswith("auction_") and column in {"high_yield", "high_discount_margin", "interest_rate", "spread"}:
        return "잠긴 auction fixture는 Bill만 포함. Note/Bond·TIPS·FRN 값을 생성하지 않음"
    return "공식 offline fixture가 해당 선택 필드를 제공하지 않음. 추정·대체값 미생성"


def column_coverage(engine: Engine, tables: list[str], spec_columns: dict[tuple[str, str], dict[str, str]]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    with engine.connect() as connection:
        for table in tables:
            total = int(connection.execute(text(f'SELECT count(*) FROM ust."{table}"')).scalar_one())
            for column in _column_metadata_for_object(engine, table, spec_columns):
                quoted = f'"{column["name"]}"'
                non_null = int(connection.execute(text(
                    f'SELECT count({quoted}) FROM ust."{table}"'
                )).scalar_one())
                numeric = bool(re.match(r"^(numeric|smallint|integer|bigint|real|double precision)", column["type"]))
                zero_count = None
                if numeric:
                    zero_count = int(connection.execute(text(
                        f'SELECT count(*) FROM ust."{table}" WHERE {quoted}=0'
                    )).scalar_one())
                example = connection.execute(text(
                    f'SELECT {quoted} FROM ust."{table}" WHERE {quoted} IS NOT NULL LIMIT 1'
                )).scalar_one_or_none()
                rows.append([
                    table, column["name"], column["description"], column["type"], column["unit"],
                    "YES" if column["nullable"] else "NO", total, non_null, zero_count,
                    excel_value(example), _all_null_reason(table, column["name"]) if total and non_null == 0 else None,
                ])
    return rows


def lineage_rows(engine: Engine) -> list[list[Any]]:
    rows: list[list[Any]] = []
    with engine.connect() as connection:
        auctions = connection.execute(text("""
            SELECT e.auction_event_id, ar.revision_number, s.source_snapshot_id, s.source_request_id,
                   s.source_url, s.content_sha256, s.raw_storage_uri, l.source_locator, l.confidence,
                   r.stop_metric_code, r.stop_value
            FROM ust.auction_event e
            JOIN ust.auction_revision ar ON ar.auction_event_id=e.auction_event_id AND ar.is_current
            JOIN ust.source_snapshot s ON s.source_snapshot_id=e.current_snapshot_id
            LEFT JOIN ust.normalized_record_lineage l ON l.entity_type='auction_event'
                 AND l.entity_id=e.auction_event_id AND l.source_snapshot_id=s.source_snapshot_id
            LEFT JOIN ust.auction_result r ON r.auction_event_id=e.auction_event_id
            ORDER BY e.auction_date,e.cusip
        """)).all()
        for item in auctions:
            rows.append([
                "입찰 Stop·응찰배수·배정", "v_auction_dashboard.stop_value/bid_to_cover_ratio",
                item[0], f"auction_revision:{item[1]}", item[2], item[3], item[4], item[5], item[6], item[7],
                "OFFICIAL_ACTUAL", "US_TREASURY_FISCAL_DATA", item[8],
                "ANNOUNCED" if item[10] is None else "RESULT_AVAILABLE",
            ])

        fact_specs = {
            "qra_borrowing_estimate": ("qra_borrowing_estimate_id", "QRA 차입 전망", "v_qra_comparison.borrowing_amount_usd"),
            "qra_auction_size": ("qra_auction_size_id", "QRA 입찰 규모/TBAC", "v_qra_auction_size_comparison.auction_size_usd"),
            "qra_supply": ("qra_supply_id", "QRA Gross/Maturing/Net", "v_qra_supply_monitor.net_issuance_usd"),
            "qra_financing_mix": ("qra_financing_mix_id", "QRA financing mix", "qra_financing_mix.implied_change_in_bills_usd"),
            "qra_tga_path": ("qra_tga_path_id", "QRA TGA 경로", "v_qra_tga_path.balance_usd"),
            "qra_guidance": ("qra_guidance_id", "QRA guidance", "qra_guidance.guidance_status"),
            "qra_tentative_auction": ("qra_tentative_auction_id", "잠정 입찰 일정", "qra_tentative_auction.auction_date"),
            "qra_buyback_operation": ("qra_buyback_operation_id", "잠정 buyback 일정", "qra_buyback_operation.operation_date"),
            "qra_buyback_policy": ("qra_buyback_policy_id", "buyback 정책 상한", "qra_buyback_policy.maximum_amount_usd"),
        }
        for table, (id_column, screen, view_column) in fact_specs.items():
            facts = connection.execute(text(f"""
                SELECT f.{id_column}, v.version_number, s.source_snapshot_id, s.source_request_id,
                       s.source_url, s.content_sha256, s.raw_storage_uri, f.source_locator, f.value_class,
                       f.source_authority, v.parse_status
                FROM ust.{table} f
                JOIN ust.qra_document_version v USING(qra_document_version_id)
                JOIN ust.source_snapshot s USING(source_snapshot_id)
                ORDER BY f.{id_column}
            """)).all()
            for item in facts:
                locator = item[7] or {}
                confidence = locator.get("extraction_confidence") if isinstance(locator, dict) else None
                rows.append([
                    screen, view_column, item[0], f"qra_document_version:{item[1]}", item[2], item[3],
                    item[4], item[5], item[6], locator, item[8], item[9], confidence, item[10],
                ])

        dealer = connection.execute(text("""
            SELECT v.qra_dealer_survey_value_id, d.version_number, ss.source_snapshot_id, ss.source_request_id,
                   ss.source_url, ss.content_sha256, ss.raw_storage_uri, v.source_locator, s.value_class, s.source_authority,
                   s.validation_status
            FROM ust.qra_dealer_survey_value v
            JOIN ust.qra_dealer_survey s USING(qra_dealer_survey_id)
            JOIN ust.qra_document_version d USING(qra_document_version_id)
            JOIN ust.source_snapshot ss USING(source_snapshot_id)
            ORDER BY v.qra_dealer_survey_value_id
        """)).all()
        for item in dealer:
            locator = item[7] or {}
            rows.append([
                "Dealer expected 규모", "v_qra_dealer_outlook.expected_auction_size_usd",
                item[0], f"qra_document_version:{item[1]}", item[2], item[3], item[4], item[5], item[6],
                locator, item[8], item[9], locator.get("extraction_confidence"), item[10],
            ])

        derived = connection.execute(text("""
            SELECT r.derived_metric_result_id, v.version_number, s.source_snapshot_id, s.source_request_id,
                   s.source_url, s.content_sha256, s.raw_storage_uri, r.input_values, r.value_class, m.source_authority,
                   m.formula
            FROM ust.derived_metric_result r
            JOIN ust.derived_metric_method m USING(derived_metric_method_id)
            JOIN ust.qra_document_version v ON v.qra_document_version_id=r.source_document_version_id
            JOIN ust.source_snapshot s USING(source_snapshot_id)
            ORDER BY r.derived_metric_result_id
        """)).all()
        for item in derived:
            rows.append([
                "ui_proxy_v1 10Y-equivalent", "v_qra_supply_monitor.dv01_10y_equivalent_proxy_usd",
                item[0], f"qra_document_version:{item[1]}", item[2], item[3], item[4], item[5], item[6],
                item[7], item[8], item[9], None, f"DERIVED: {item[10]}",
            ])
    return rows


def validation_rows(engine: Engine, table_counts: dict[str, int], view_counts: dict[str, int]) -> list[list[Any]]:
    repository = Repository(engine=engine)
    rows: list[list[Any]] = []
    for name, count in sorted(demo_validation(repository).items()):
        rows.append([name, count, "PASS" if count == 0 else "FAIL", None, None, None, "expected result count = 0"])
    with engine.connect() as connection:
        checkpoint = connection.execute(text("""
            SELECT checkpoint FROM ust.ingestion_run
            WHERE source_name='DEMO_FIXTURE' AND status='SUCCESS'
            ORDER BY started_at DESC LIMIT 1
        """)).scalar_one_or_none() or {}
    for name, value in sorted((checkpoint.get("idempotency") or {}).items()):
        if name.endswith("_before") or name.endswith("_after"):
            continue
        status = "PASS"
        if "change" in name and value != 0:
            status = "FAIL"
        if "increase" in name and value != 2:
            status = "FAIL"
        rows.append([f"idempotency.{name}", value if isinstance(value, int) else None, status, None, None, value, stable_json(value)])
    for table, count in table_counts.items():
        rows.append([f"table_nonempty.{table}", 0 if count > 0 else 1, "PASS" if count > 0 else "FAIL", None, count, None, f"rows={count}"])
    for view, count in view_counts.items():
        rows.append([f"view_select.{view}", 0, "PASS", None, count, None, f"rows={count}"])
    return rows


def build_workbook(engine: Engine) -> tuple[Workbook, dict[str, str], dict[str, str], dict[str, int], dict[str, int]]:
    root = repository_root()
    manifest, manifest_audit = verify_manifest(root)
    assert_safe_database_state(engine, allow_empty=False)
    info = server_info(engine)
    tables, views = catalog_objects(engine)
    table_counts, view_counts = object_counts(Repository(engine=engine))
    if not all(table_counts.values()):
        raise ValidationFailure("all catalog tables must have demo rows before export")
    if not all(view_counts.values()):
        raise ValidationFailure("all catalog views must be queryable and non-empty before export")
    spec_purposes, spec_columns = parse_db_spec(root)
    used = set(SPECIAL_SHEETS)
    table_sheets = {name: safe_sheet_name("T", name, used) for name in tables}
    view_sheets = {name: safe_sheet_name("V", name, used) for name in views}

    workbook = Workbook()
    workbook.remove(workbook.active)

    readme = workbook.create_sheet("00_README")
    now = datetime.now(UTC)
    readme_rows = [
        ["generated_at_utc", now.isoformat()],
        ["generated_at_kst", now.astimezone(ZoneInfo("Asia/Seoul")).isoformat()],
        ["database", info["database"]],
        ["postgresql_server", info["server_version"]],
        ["alembic_version", migration_version(engine)],
        ["demo_source", f"DEMO_FIXTURE: MANIFEST.json의 공식 offline fixture {len(manifest_audit)}개"],
        ["OFFICIAL_ACTUAL", "공식 실적 또는 확정 입찰결과"],
        ["OFFICIAL_ESTIMATE", "Treasury 공식 전망·잠정 일정"],
        ["TBAC_RECOMMENDATION", "TBAC 권고. Treasury 결정과 별도"],
        ["PRIMARY_DEALER_SURVEY", "Primary Dealer 응답 통계. current 실적이 아님"],
        ["DERIVED_UI_PROXY", "ui_proxy_v1의 Net×tenor weight. 공식 Treasury DV01 지표가 아님"],
        ["NULL_vs_zero", "NULL은 미제공·비적용·미연결. 0은 원천에 실제 0이 있는 값이며 별도 셀로 보존"],
        ["market_data_gap", "WI, Tail, 실시간 시장금리는 지정 공식 fixture에 없어 미제공. Stop으로 대체하지 않음"],
        ["dealer_gap", "current_auction_size_usd, calculated_pct_change, expected_increase_timing은 field-level 근거가 없어 빈 셀"],
        ["auction_fixture_gap", "공식 auction fixture는 Bill 4건뿐. CMB/Note/Bond/TIPS/FRN Stop과 동일 CUSIP 두 번째 사건을 생성하지 않음"],
        ["bid_to_cover_display", "bid_to_cover_ratio는 배수 원값. 화면 %는 v_auction_dashboard.bid_to_cover_display_pct에서 ×100"],
        ["UI_SHA256_expected", "D5A19347BAF26EA48FFF6E159E7B0992FEA67481FBD146804F9FF2B62E5F1B4E"],
    ]
    _write_matrix(readme, ["item", "description"], readme_rows)

    table_index_rows: list[list[Any]] = []
    for table in tables:
        pk, uniques = constraint_metadata(engine, table)
        update_mode, retention = update_policy(table)
        table_index_rows.append([
            table, TABLE_PURPOSES.get(table, spec_purposes.get(table, table)), table_counts[table],
            table_sheets[table], ", ".join(pk), " | ".join(uniques), update_mode, retention,
        ])
    _write_matrix(
        workbook.create_sheet("01_TABLE_INDEX"),
        ["table_name", "purpose_ko", "row_count", "sheet_name", "primary_key", "business_key_or_unique", "update_mode", "retention_policy"],
        table_index_rows,
    )

    view_index_rows = [
        [view, VIEW_METADATA.get(view, (view, "DB view 계산", "NULL 보존"))[0], view_counts[view], view_sheets[view],
         VIEW_METADATA.get(view, (view, "DB view 계산", "NULL 보존"))[1], VIEW_METADATA.get(view, (view, "DB view 계산", "NULL 보존"))[2]]
        for view in views
    ]
    _write_matrix(
        workbook.create_sheet("02_VIEW_INDEX"),
        ["view_name", "screen_connection", "row_count", "sheet_name", "major_calculation", "missing_value_handling"],
        view_index_rows,
    )

    _write_matrix(
        workbook.create_sheet("03_COLUMN_COVERAGE"),
        ["table_name", "column_name", "description_ko", "postgresql_type", "unit", "nullable", "total_rows", "non_null_count", "actual_zero_count", "example_value", "all_null_reason"],
        column_coverage(engine, tables, spec_columns),
    )
    _write_matrix(
        workbook.create_sheet("04_LINEAGE_TRACE"),
        ["screen_element", "view_column", "fact_record_id", "revision_or_document_version", "source_snapshot_id", "source_request_id", "source_url", "raw_content_sha256", "raw_storage_uri", "source_locator", "value_class", "source_authority", "confidence", "validation_status"],
        lineage_rows(engine),
    )
    _write_matrix(
        workbook.create_sheet("05_VALIDATION"),
        ["validation_rule", "result_count", "status", "before_value", "after_value", "change", "details"],
        validation_rows(engine, table_counts, view_counts),
    )

    for table in tables:
        pk, _ = constraint_metadata(engine, table)
        columns, rows = _query_rows(engine, table, pk)
        metadata = _column_metadata_for_object(engine, table, spec_columns)
        _write_matrix(workbook.create_sheet(table_sheets[table]), columns, [list(row) for row in rows], metadata)
    for view in views:
        columns, rows = _query_rows(engine, view)
        metadata = _column_metadata_for_object(engine, view, spec_columns)
        _write_matrix(workbook.create_sheet(view_sheets[view]), columns, [list(row) for row in rows], metadata)
    return workbook, table_sheets, view_sheets, table_counts, view_counts


def validate_saved_workbook(
    path: Path,
    engine: Engine,
    table_sheets: dict[str, str],
    view_sheets: dict[str, str],
    table_counts: dict[str, int],
    view_counts: dict[str, int],
) -> dict[str, Any]:
    loaded = load_workbook(path, read_only=False, data_only=False)
    try:
        if len(loaded.sheetnames) != len(set(name.casefold() for name in loaded.sheetnames)):
            raise ValidationFailure("duplicate workbook sheet names")
        if any(len(name) > 31 or FORBIDDEN_SHEET_CHARS.search(name) for name in loaded.sheetnames):
            raise ValidationFailure("invalid workbook sheet name")
        expected = set(SPECIAL_SHEETS) | set(table_sheets.values()) | set(view_sheets.values())
        if set(loaded.sheetnames) != expected:
            raise ValidationFailure("workbook object/sheet coverage mismatch")
        for object_name, sheet_name in {**table_sheets, **view_sheets}.items():
            ws = loaded[sheet_name]
            db_columns = [item["name"] for item in column_catalog(engine, object_name)]
            headers = [ws.cell(1, index).value for index in range(1, ws.max_column + 1)]
            if headers != db_columns:
                raise ValidationFailure(f"Excel headers differ from DB columns: {object_name}")
            expected_rows = table_counts.get(object_name, view_counts.get(object_name, 0))
            if ws.max_row - 1 != expected_rows:
                raise ValidationFailure(f"Excel row count differs from DB: {object_name}")
            if ws.freeze_panes != "A2" or ws.auto_filter.ref is None:
                raise ValidationFailure(f"freeze/filter missing: {sheet_name}")
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type == "f":
                        raise ValidationFailure(f"unexpected formula cell: {sheet_name}!{cell.coordinate}")
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if any("externalLinks" in name or "vbaProject" in name for name in names):
                raise ValidationFailure("workbook contains external links or macros")
        return {
            "status": "PASS",
            "sheets": len(loaded.sheetnames),
            "tables": len(table_sheets),
            "views": len(view_sheets),
            "table_rows": sum(table_counts.values()),
            "view_rows": sum(view_counts.values()),
            "reopened": True,
            "formula_cells": 0,
            "external_links": 0,
            "macros": 0,
        }
    finally:
        loaded.close()


def export_workbook(output: Path, *, database_url: str | None = None) -> dict[str, Any]:
    url = database_url or demo_database_url()
    engine = create_demo_engine(url)
    try:
        workbook, table_sheets, view_sheets, table_counts, view_counts = build_workbook(engine)
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
        report = validate_saved_workbook(output, engine, table_sheets, view_sheets, table_counts, view_counts)
        report.update(output=str(output.resolve()), bytes=output.stat().st_size)
        report_path = repository_root() / "work" / "phase-13" / "workbook-validation.json"
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report
    finally:
        engine.dispose()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Export every UST demo table and view to one XLSX workbook")
    parser.add_argument("--output", type=Path, default=Path("artifacts/UST_Auction_QRA_DB_Demo.xlsx"))
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        report = export_workbook(args.output)
        print(json.dumps(report, ensure_ascii=False, sort_keys=True, indent=2, default=str))
        return 0
    except (ValidationFailure, DemoSafetyError) as exc:
        print(f"VALIDATION FAILED: {exc}", file=sys.stderr)
        return 2
    except Exception as exc:
        print(f"EXECUTION FAILED: {exc}", file=sys.stderr)
        return 1
