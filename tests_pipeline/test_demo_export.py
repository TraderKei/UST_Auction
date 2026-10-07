from __future__ import annotations

import os
import re
from datetime import date, time
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.exc import IntegrityError

from ust_pipeline.demo.common import DemoSafetyError, demo_database_url, migration_version
from ust_pipeline.demo.loader import load_demo
from ust_pipeline.demo.workbook import (
    FORBIDDEN_SHEET_CHARS,
    SPECIAL_SHEETS,
    column_catalog,
    constraint_metadata,
    excel_value,
    export_workbook,
    safe_sheet_name,
    safe_source_text,
)


@pytest.fixture(scope="module")
def demo_url() -> str:
    value = os.getenv("UST_DEMO_DATABASE_URL")
    if not value:
        pytest.skip("UST_DEMO_DATABASE_URL must point to a dedicated demo/test PostgreSQL 16 database")
    return value


@pytest.fixture(scope="module")
def loaded_demo(demo_url: str) -> dict:
    return load_demo(database_url=demo_url)


def test_demo_url_safety_rejects_missing_unapproved_and_production_reuse():
    with pytest.raises(DemoSafetyError):
        demo_database_url({})
    with pytest.raises(DemoSafetyError):
        demo_database_url({"UST_DEMO_DATABASE_URL": "postgresql+psycopg://u@localhost/treasury"})
    same = "postgresql+psycopg://u@localhost/treasury_demo"
    with pytest.raises(DemoSafetyError):
        demo_database_url({"UST_DEMO_DATABASE_URL": same, "UST_DATABASE_URL": same})
    assert demo_database_url({"UST_DEMO_DATABASE_URL": same}) == same


def test_excel_serialization_sheet_names_and_formula_injection():
    assert excel_value(Decimal("123456789012345.678900000")) == "123456789012345.678900000"
    assert isinstance(excel_value(date(2026, 8, 5)), date)
    assert isinstance(excel_value(time(13, 30)), time)
    assert safe_source_text("=2+2") == "'=2+2"
    assert safe_source_text("@SUM(A1:A2)") == "'@SUM(A1:A2)"
    used = set(SPECIAL_SHEETS)
    first = safe_sheet_name("T", "derived_metric_method_parameter", used)
    second = safe_sheet_name("T", "derived_metric_method_parameter", used)
    assert len(first) <= 31 and len(second) <= 31 and first != second
    assert not FORBIDDEN_SHEET_CHARS.search(first)


@pytest.mark.integration
def test_empty_database_migrates_and_all_catalog_objects_are_populated(loaded_demo, demo_url):
    assert loaded_demo["status"] == "SUCCESS"
    assert len(loaded_demo["tables"]) == 27
    assert len(loaded_demo["views"]) == 8
    assert all(count >= 1 for count in loaded_demo["tables"].values())
    assert all(count >= 1 for count in loaded_demo["views"].values())
    engine = create_engine(demo_url)
    try:
        assert migration_version(engine) == "20261007_0002"
        assert int(engine.connect().execute(text("SHOW server_version_num")).scalar_one()) >= 160000
    finally:
        engine.dispose()


@pytest.mark.integration
def test_repeat_load_does_not_increase_current_revision_version_or_lineage(loaded_demo, demo_url):
    engine = create_engine(demo_url)
    try:
        with engine.connect() as connection:
            before = connection.execute(text("""
                SELECT (SELECT count(*) FROM ust.auction_revision),
                       (SELECT count(*) FROM ust.qra_document_version),
                       (SELECT count(*) FROM ust.normalized_record_lineage)
            """)).one()
        repeated = load_demo(database_url=demo_url)
        with engine.connect() as connection:
            after = connection.execute(text("""
                SELECT (SELECT count(*) FROM ust.auction_revision),
                       (SELECT count(*) FROM ust.qra_document_version),
                       (SELECT count(*) FROM ust.normalized_record_lineage)
            """)).one()
        assert repeated["repeat"] is True
        assert before == after
        assert repeated["idempotency"]["auction_revision_change"] == 0
        assert repeated["idempotency"]["qra_document_version_change"] == 0
        assert repeated["idempotency"]["lineage_change"] == 0
    finally:
        engine.dispose()


@pytest.mark.integration
def test_aba_quarantine_status_zero_and_official_gaps(loaded_demo, demo_url):
    engine = create_engine(demo_url)
    try:
        with engine.connect() as connection:
            revisions = connection.execute(text("""
                SELECT revision_number,content_hash FROM ust.auction_revision r
                JOIN ust.auction_event e USING(auction_event_id)
                WHERE e.cusip='912797UM7' ORDER BY revision_number
            """)).all()
            assert len(revisions) == 3
            assert revisions[0][1] == revisions[2][1] and revisions[0][1] != revisions[1][1]
            qra_versions = connection.execute(text("""
                SELECT version_number,content_sha256 FROM ust.qra_document_version v
                JOIN ust.qra_document d USING(qra_document_id)
                WHERE d.document_type='QRA_INDEX_HTML' ORDER BY version_number
            """)).all()
            assert len(qra_versions) == 3
            assert qra_versions[0][1] == qra_versions[2][1] and qra_versions[0][1] != qra_versions[1][1]
            quarantine = connection.execute(text("""
                SELECT v.parse_status,
                       (SELECT count(*) FROM ust.qra_borrowing_estimate b WHERE b.qra_document_version_id<>v.qra_document_version_id)
                FROM ust.qra_document_version v JOIN ust.qra_document d USING(qra_document_id)
                WHERE d.document_type='FINANCING_ESTIMATES_HTML' AND v.is_current
            """)).one()
            assert quarantine[0] == "QUARANTINED" and quarantine[1] == 3
            statuses = {row[0] for row in connection.execute(text("SELECT event_status FROM ust.v_auction_dashboard"))}
            assert statuses == {"ANNOUNCED", "RESULT_AVAILABLE"}
            assert connection.execute(text("""
                SELECT count(*) FROM ust.auction_bidder_allocation
                WHERE fima_noncomp_accepted_usd=0 OR fima_noncomp_tendered_usd=0
            """)).scalar_one() >= 1
            assert connection.execute(text("""
                SELECT count(*) FROM ust.qra_dealer_survey_value
                WHERE current_auction_size_usd IS NOT NULL OR calculated_pct_change IS NOT NULL
            """)).scalar_one() == 0
            assert connection.execute(text("""
                SELECT count(*) FROM information_schema.columns WHERE table_schema='ust' AND
                  (column_name LIKE 'wi_%' OR column_name LIKE 'tail%' OR column_name LIKE 'market_rate%')
            """)).scalar_one() == 0
            gaps = {row[0] for row in connection.execute(text("""
                SELECT rule_code FROM ust.data_quality_issue
                WHERE rule_code IN ('DEMO_FIXTURE_STOP_COVERAGE_GAP','DEMO_FIXTURE_SAME_CUSIP_GAP')
            """))}
            assert gaps == {"DEMO_FIXTURE_STOP_COVERAGE_GAP", "DEMO_FIXTURE_SAME_CUSIP_GAP"}
    finally:
        engine.dispose()


@pytest.mark.integration
def test_fk_unique_and_check_constraints_remain_enforced(loaded_demo, demo_url):
    engine = create_engine(demo_url)
    try:
        with engine.connect() as connection:
            constraint_count = connection.execute(text("""
                SELECT count(*) FROM pg_constraint c JOIN pg_namespace n ON n.oid=c.connamespace
                WHERE n.nspname='ust' AND c.contype IN ('p','f','u','c')
            """)).scalar_one()
            assert constraint_count >= 100
            with pytest.raises(IntegrityError):
                connection.execute(text(
                    "INSERT INTO ust.qra_refunding(refunding_year,refunding_month,source_page_url) VALUES(2026,3,'DEMO_FIXTURE invalid')"
                ))
            connection.rollback()
    finally:
        engine.dispose()


@pytest.mark.integration
def test_workbook_reopens_and_matches_database_types_lineage_and_counts(loaded_demo, demo_url, tmp_path):
    output = tmp_path / "demo.xlsx"
    report = export_workbook(output, database_url=demo_url)
    assert report["status"] == "PASS" and report["sheets"] == 41
    engine = create_engine(demo_url)
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        index = workbook["01_TABLE_INDEX"]
        mapping = {index.cell(row, 1).value: index.cell(row, 4).value for row in range(2, index.max_row + 1)}
        view_index = workbook["02_VIEW_INDEX"]
        view_mapping = {view_index.cell(row, 1).value: view_index.cell(row, 4).value for row in range(2, view_index.max_row + 1)}
        assert len(mapping) == 27 and len(view_mapping) == 8
        for object_name, sheet_name in {**mapping, **view_mapping}.items():
            expected_headers = [item["name"] for item in column_catalog(engine, object_name)]
            ws = workbook[sheet_name]
            assert [ws.cell(1, col).value for col in range(1, ws.max_column + 1)] == expected_headers
            assert ws.freeze_panes == "A2" and ws.auto_filter.ref
        result_ws = workbook[mapping["auction_result"]]
        ratio_col = [cell.value for cell in result_ws[1]].index("bid_to_cover_ratio") + 1
        assert all(
            result_ws.cell(row, ratio_col).value is None or isinstance(result_ws.cell(row, ratio_col).value, str)
            for row in range(2, result_ws.max_row + 1)
        )
        event_ws = workbook[mapping["auction_event"]]
        auction_date_col = [cell.value for cell in event_ws[1]].index("auction_date") + 1
        time_col = [cell.value for cell in event_ws[1]].index("closing_time_comp_et") + 1
        assert isinstance(event_ws.cell(2, auction_date_col).value, date)
        assert isinstance(event_ws.cell(2, time_col).value, time)
        run_ws = workbook[mapping["ingestion_run"]]
        started_col = [cell.value for cell in run_ws[1]].index("started_at") + 1
        assert re.search(r"[+-]\d\d:\d\d$", run_ws.cell(2, started_col).value)
        assert not any(cell.data_type == "f" for ws in workbook.worksheets for row in ws.iter_rows() for cell in row)
        lineage = workbook["04_LINEAGE_TRACE"]
        headers = [cell.value for cell in lineage[1]]
        assert "raw_content_sha256" in headers and "source_url" in headers and "source_locator" in headers
        hash_col = headers.index("raw_content_sha256") + 1
        assert all(re.fullmatch(r"[0-9a-f]{64}", lineage.cell(row, hash_col).value) for row in range(2, lineage.max_row + 1))
        dealer_ws = workbook[mapping["qra_dealer_survey_value"]]
        current_col = [cell.value for cell in dealer_ws[1]].index("current_auction_size_usd") + 1
        delta_col = [cell.value for cell in dealer_ws[1]].index("calculated_pct_change") + 1
        assert all(dealer_ws.cell(row, current_col).value is None for row in range(2, dealer_ws.max_row + 1))
        assert all(dealer_ws.cell(row, delta_col).value is None for row in range(2, dealer_ws.max_row + 1))
    finally:
        workbook.close()
        engine.dispose()
